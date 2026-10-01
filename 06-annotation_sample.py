# -*- coding: utf-8 -*-
"""
06 - Stratified sample for manual annotation (two annotators, Word files)
=========================================================================
Why: step 05 showed that the complex > simple difference comes mainly from
RARE predicates (fewer than 10 overt objects), where parsing errors are most
likely. Before any model, we must know what the "omitted" tokens really are.

Sample (500 items):
    omitted tokens, 100 per cell of
        predicate type  (complex / simple, CP-strict)  x
        frequency band  (reliable = >= 10 overt objects / rare)
    + 100 tokens WITH an overt object (control: how often is a visible
      object itself a parsing error?)
Items are shuffled, so every batch mixes all strata, and the strata are
hidden from the annotators (blind annotation).

Output: 5 Word files of 100 items for EACH annotator (10 files in total):
    annotation/annotator1/06_batch1.docx ... 06_batch5.docx
    annotation/annotator2/06_batch1.docx ... 06_batch5.docx
    annotation/06_sample_key.csv          (hidden strata, for step 07 only)

Each file starts with the guide, then a right-to-left table:
    شماره | محمول | جمله (verb in bold) | برچسب | توضیح
The annotator writes ONE code in the «برچسب» column (Persian or Latin):
    آ  (F)  حذف آزاد / بافت‌آزاد
    ض  (G)  مفعول به‌صورت مضاف‌الیهِ جزء غیرفعلی: «طلبِ کرامت کردند»
    ق  (C)  حذف به قرینه
    خ  (E)  خطای برچسب‌گذاری: مفعول در جمله هست
    غ  (N)  کاربرد غیرمتعدی یا ساخت دیگر
    ؟  (?)  مطمئن نیستم

Requires: pip install python-docx
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from conllu_reader import read_all

SEED = 2026
PER_CELL = 100
N_CONTROL = 100
BATCH_SIZE = 100
FONT = "Tahoma"          # available on every Windows machine, supports Persian
OUT = Path("annotation")
OUT.mkdir(exist_ok=True)

# ---------------------------------------------------------------------------
# 1. Sample
# ---------------------------------------------------------------------------
tok = pd.read_csv("outputs/tables/05_analysis_tokens.csv", encoding="utf-8-sig")
tok["reliable"] = tok["reliable"].fillna(False).astype(bool)

parts = []
om = tok[tok["omitted"]]
for cp in [True, False]:
    for rel in [True, False]:
        cell = om[(om["cp_strict"] == cp) & (om["reliable"] == rel)]
        k = min(PER_CELL, len(cell))
        smp = cell.sample(n=k, random_state=SEED).copy()
        smp["stratum"] = f"{'complex' if cp else 'simple'}_{'reliable' if rel else 'rare'}_omitted"
        parts.append(smp)
        print(f"{smp['stratum'].iloc[0]:32s} available {len(cell):5d}   sampled {k}")
ctrl = tok[~tok["omitted"]].sample(n=N_CONTROL, random_state=SEED).copy()
ctrl["stratum"] = "control_with_object"
parts.append(ctrl)
print(f"{'control_with_object':32s} sampled {N_CONTROL}")

sample = pd.concat(parts).sample(frac=1, random_state=SEED).reset_index(drop=True)
sample["item"] = range(1, len(sample) + 1)
sample["batch"] = (sample["item"] - 1) // BATCH_SIZE + 1

# words of each sentence, to rebuild it with the verb in bold
words = {}
for s in read_all():
    words[(s["file"], s["sent_id"])] = [(t["id"], t["form"]) for t in s["tokens"]]


# ---------------------------------------------------------------------------
# 2. Right-to-left helpers for Word
# ---------------------------------------------------------------------------
def rtl_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    return p


def add_run(p, text, size=11, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(OxmlElement("w:rtl"))
    sz = OxmlElement("w:szCs")
    sz.set(qn("w:val"), str(size * 2))
    rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    return r


def para(doc_or_cell, text="", size=11, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = rtl_paragraph(doc_or_cell.add_paragraph(), align)
    if text:
        add_run(p, text, size, bold)
    return p


def rtl_table(table, widths=None):
    tblpr = table._tbl.tblPr
    tblpr.append(OxmlElement("w:bidiVisual"))
    if widths:
        table.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w.twips)))


def set_cell_text(cell, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = cell.paragraphs[0]
    rtl_paragraph(p, align)
    add_run(p, text, size, bold)


# ---------------------------------------------------------------------------
# 3. Guide (Persian)
# ---------------------------------------------------------------------------
GUIDE = [
    ("آ", "F", "حذف آزاد (بافت‌آزاد)",
     "مفعول نیامده و مرجع آن در جمله یا بافت ذکر نشده است.",
     "«دائم داره می‌خوره.»  «روی پوست درختان حکاکی نکنید.»"),
    ("ض", "G", "مفعول به‌صورت مضاف‌الیه",
     "مفعول آمده، ولی نه جدا: به‌صورت مضاف‌الیهِ جزء غیرفعلی (با کسرهٔ اضافه).",
     "«طلبِ کرامت کردند»  «طلبِ باران می‌کردند»  (= کرامت/باران را طلب کردند)"),
    ("ق", "C", "حذف به قرینه",
     "مفعول نیامده، ولی مرجع آن در همین جمله یا جملهٔ قبل هست؛ مثلاً مفعول مشترک با فعل هم‌پایه.",
     "«کتاب را باز کنند و بخوانند.»"),
    ("خ", "E", "خطای برچسب‌گذاری",
     "مفعول صریح در جمله هست، ولی پیکره آن را مفعول نشناخته است.",
     "«یک جفت پتو خریده بود.» (پتو مفعول است)"),
    ("غ", "N", "غیرمتعدی یا ساخت دیگر",
     "فعل در این جمله اصلاً مفعول صریح نمی‌طلبد: معنای دیگر، اصطلاح، فعل نمودی یا متمم حرف‌اضافه‌ای.",
     "«شروع کرد به خواندن.»  «به درد می‌خورد.»  «به زمین خورد.»"),
    ("؟", "?", "مطمئن نیستم",
     "در ستون توضیح بنویسید چرا.",
     ""),
]


GUIDE_WIDTHS = [Cm(1.2), Cm(1.3), Cm(3.8), Cm(9.0), Cm(9.4)]


def write_guide(doc, batch, annotator):
    para(doc, f"برچسب‌گذاری حذف مفعول — برچسب‌زن {annotator} — بستهٔ {batch} از 5",
         size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "در هر جمله، فعل مورد نظر پررنگ و قرمز است. برای هر سطر فقط یک کد در "
              "ستون «برچسب» بنویسید. کد فارسی یا لاتین هر دو پذیرفته است.", size=10)
    para(doc, "در سطرهایی که فعل مفعول دارد، اگر مفعول درست است «درست» بنویسید؛ "
              "اگر مفعول در واقع وجود ندارد، همان کدهای بالا را به کار ببرید.", size=10)
    para(doc, "مستقل از برچسب‌زن دیگر کار کنید و با او مشورت نکنید.", size=10, bold=True)

    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    rtl_table(t, GUIDE_WIDTHS)
    for c, h, w in zip(t.rows[0].cells, ["کد", "لاتین", "معنی", "توضیح", "مثال"], GUIDE_WIDTHS):
        set_cell_text(c, h, 10, True, WD_ALIGN_PARAGRAPH.CENTER)
        c.width = w
    for fa, la, name, desc, ex in GUIDE:
        row = t.add_row().cells
        for c, w in zip(row, GUIDE_WIDTHS):
            c.width = w
        set_cell_text(row[0], fa, 12, True, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row[1], la, 10, False, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row[2], name, 10, True)
        set_cell_text(row[3], desc, 9)
        set_cell_text(row[4], ex, 9)
    para(doc, "")


# ---------------------------------------------------------------------------
# 4. Items table
# ---------------------------------------------------------------------------
WIDTHS = [Cm(1.5), Cm(3.2), Cm(13.8), Cm(1.8), Cm(4.4)]


def write_items(doc, items):
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    rtl_table(t, WIDTHS)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))   # repeat on each page
    for c, h, w in zip(t.rows[0].cells, ["شماره", "محمول", "جمله", "برچسب", "توضیح"], WIDTHS):
        set_cell_text(c, h, 10, True, WD_ALIGN_PARAGRAPH.CENTER)
        c.width = w
    for _, r in items.iterrows():
        cells = t.add_row().cells
        for c, w in zip(cells, WIDTHS):
            c.width = w
        set_cell_text(cells[0], str(r["item"]), 9, False, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], r["predicate"], 10, True)
        p = rtl_paragraph(cells[2].paragraphs[0])
        for i, w in words[(r["file"], r["sent_id"])]:
            if i == r["tok_id"]:
                add_run(p, w, 11, bold=True, color=(192, 0, 0))
            else:
                add_run(p, w, 11)
            add_run(p, " ", 11)
        rtl_paragraph(cells[3].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        rtl_paragraph(cells[4].paragraphs[0])


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    return doc


for annotator in [1, 2]:
    folder = OUT / f"annotator{annotator}"
    folder.mkdir(exist_ok=True)
    for b, items in sample.groupby("batch"):
        doc = new_document()
        write_guide(doc, b, annotator)
        write_items(doc, items)
        target = folder / f"06_batch{b}.docx"
        if target.exists():
            print(f"  kept existing {target} (delete it to regenerate)")
            continue
        doc.save(target)

sample[["item", "batch", "stratum", "file", "sent_id", "tok_id", "predicate",
        "cp_strict", "cp_broad", "light_verb", "reliable", "has_obj",
        "conjunct_has_obj", "imperfective", "imperative"]].to_csv(
    OUT / "06_sample_key.csv", index=False, encoding="utf-8-sig")

print(f"\nTotal items: {len(sample)} in {sample['batch'].nunique()} batches of {BATCH_SIZE}")
print("Saved: annotation/annotator1/06_batch1-5.docx, annotation/annotator2/06_batch1-5.docx,")
print("       annotation/06_sample_key.csv")
